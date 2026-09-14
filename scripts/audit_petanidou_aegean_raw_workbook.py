from __future__ import annotations

import hashlib
import io
import json
import re
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/results/chapter2_petanidou_aegean_raw_workbook_audit_20260915.json"
SOURCE_REPO = "JoseBSL/EuPPollNet"
TAG = "v1.3.0"
RAW_BLOB_SHA = "f238775bc6568a24a727a62524f535adad4de8f2"
RAW_URL = f"https://raw.githubusercontent.com/{SOURCE_REPO}/{TAG}/Data/1_Raw_data/51_Petanidou/Raw.xlsx"


def get_bytes(url: str, timeout: int = 120) -> bytes:
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": "izu-core-source-audit/1.0",
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*;q=0.8",
        },
    )
    with urllib.request.urlopen(req, timeout=timeout) as response:
        return response.read()


def clean(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", clean(value).casefold()).strip("_")


def row_values(row: tuple[object, ...]) -> list[str]:
    return [clean(v) for v in row]


def score_header(values: list[str]) -> int:
    tokens = (
        "date", "day", "month", "year", "site", "locality", "island", "location",
        "pollinator", "visitor", "insect", "bee", "plant", "flower", "interaction",
        "visit", "effort", "minute", "duration", "area", "sampling", "round", "transect",
    )
    return sum(any(token in norm(v) for token in tokens) for v in values if v)


def detect_header(rows: list[tuple[object, ...]]) -> tuple[int | None, list[str]]:
    best_row = None
    best_values: list[str] = []
    best_score = 0
    for i, row in enumerate(rows[:40], start=1):
        values = row_values(row)
        score = score_header(values)
        if score > best_score and sum(bool(v) for v in values) >= 2:
            best_row, best_values, best_score = i, values, score
    return best_row, best_values


def roles(headers: list[str]) -> dict[str, list[str]]:
    out: dict[str, list[str]] = defaultdict(list)
    for raw in headers:
        n = norm(raw)
        if not n:
            continue
        if any(x in n for x in ("site", "locality", "island", "location", "place")):
            out["site"].append(raw)
        if any(x in n for x in ("date", "day", "month", "year", "round", "sampling")):
            out["time"].append(raw)
        if any(x in n for x in ("pollinator", "visitor", "insect", "bombus", "bee")):
            out["partner"].append(raw)
        if any(x in n for x in ("plant", "flower")):
            out["plant"].append(raw)
        if any(x in n for x in ("interaction", "visit", "frequency", "count", "abundance", "number")):
            out["interaction"].append(raw)
        if any(x in n for x in ("effort", "minute", "duration", "area", "transect")):
            out["effort"].append(raw)
    return dict(out)


def date_from_mapping(mapping: dict[str, object]) -> str:
    lowered = {norm(k): v for k, v in mapping.items()}
    if "date" in lowered and clean(lowered["date"]):
        value = lowered["date"]
        if hasattr(value, "strftime"):
            return value.strftime("%Y-%m-%d")
        return clean(value)[:10]
    day = next((lowered[k] for k in lowered if k == "day" or k.endswith("_day")), None)
    month = next((lowered[k] for k in lowered if k == "month" or k.endswith("_month")), None)
    year = next((lowered[k] for k in lowered if k == "year" or k.endswith("_year")), None)
    if all(clean(v) for v in (day, month, year)):
        try:
            return f"{int(float(year)):04d}-{int(float(month)):02d}-{int(float(day)):02d}"
        except (TypeError, ValueError):
            return ""
    return ""


def main() -> None:
    payload = get_bytes(RAW_URL)
    if not payload.startswith(b"PK\x03\x04"):
        raise RuntimeError("Raw.xlsx did not return a valid XLSX payload")
    wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    sheet_results = []
    workbook_schedule_candidates = []

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        header_row, headers = detect_header(rows)
        record: dict[str, object] = {
            "sheet": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "header_row_1_based": header_row,
            "headers": headers,
            "roles": roles(headers),
            "preview": [row_values(row)[:16] for row in rows[:15]],
        }
        if header_row is not None and headers:
            body = rows[header_row:]
            mappings = [dict(zip(headers, row)) for row in body]
            nonblank_mappings = [m for m in mappings if any(clean(v) for v in m.values())]
            date_values = [date_from_mapping(m) for m in nonblank_mappings]
            dates = sorted({v for v in date_values if v})

            role_map = roles(headers)
            site_cols = role_map.get("site", [])
            partner_cols = role_map.get("partner", [])
            interaction_cols = role_map.get("interaction", [])
            effort_cols = role_map.get("effort", [])

            per_site_dates: dict[str, set[str]] = defaultdict(set)
            dated_rows = 0
            dated_rows_without_partner = 0
            dated_rows_without_interaction_value = 0
            effort_values: dict[str, set[str]] = defaultdict(set)

            for mapping, day in zip(nonblank_mappings, date_values):
                if not day:
                    continue
                dated_rows += 1
                site = ""
                for col in site_cols:
                    if clean(mapping.get(col)):
                        site = clean(mapping.get(col))
                        break
                if not site:
                    site = ws.title
                per_site_dates[site].add(day)
                partner_present = any(clean(mapping.get(col)) for col in partner_cols)
                interaction_present = any(clean(mapping.get(col)) for col in interaction_cols)
                if partner_cols and not partner_present:
                    dated_rows_without_partner += 1
                if interaction_cols and not interaction_present:
                    dated_rows_without_interaction_value += 1
                for col in effort_cols:
                    value = clean(mapping.get(col))
                    if value:
                        effort_values[col].add(value)

            record["structure"] = {
                "distinct_dates": len(dates),
                "date_min": dates[0] if dates else None,
                "date_max": dates[-1] if dates else None,
                "dated_rows": dated_rows,
                "dated_rows_without_partner": dated_rows_without_partner,
                "dated_rows_without_interaction_value": dated_rows_without_interaction_value,
                "per_site_distinct_dates": {k: len(v) for k, v in sorted(per_site_dates.items())},
                "effort_values": {k: sorted(v)[:50] for k, v in sorted(effort_values.items())},
            }
            if dates and (partner_cols or interaction_cols or site_cols):
                workbook_schedule_candidates.append(ws.title)
        sheet_results.append(record)

    wb.close()
    result = {
        "schema_version": "1.0",
        "analysis": "chapter2_petanidou_aegean_raw_workbook_audit",
        "status": "source_structure_only_coordinates_not_opened",
        "source": {
            "repository": SOURCE_REPO,
            "tag": TAG,
            "raw_git_blob_sha": RAW_BLOB_SHA,
            "raw_url": RAW_URL,
            "raw_xlsx_bytes": len(payload),
            "raw_xlsx_sha256": hashlib.sha256(payload).hexdigest(),
        },
        "sheet_count": len(sheet_results),
        "sheets": sheet_results,
        "schedule_candidate_sheets": workbook_schedule_candidates,
        "coordinates_opened": False,
        "claim_boundary": (
            "The workbook audit opens only workbook transport, sheet structure, headers, date/site cardinalities, effort fields, and whether dated rows exist without a partner or interaction value. "
            "It does not compute interaction breadth, synchrony, outcomes, determinant rankings or route values. A site is not admitted merely because positive-interaction dates meet the six-bin floor; an outcome-independent source-native sampling frame is still required."
        ),
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({
        "raw_xlsx_sha256": result["source"]["raw_xlsx_sha256"],
        "sheet_count": result["sheet_count"],
        "schedule_candidate_sheets": workbook_schedule_candidates,
        "sheet_summaries": [
            {
                "sheet": row["sheet"],
                "headers": row["headers"],
                "structure": row.get("structure"),
            }
            for row in sheet_results
        ],
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
