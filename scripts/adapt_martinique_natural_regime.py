from __future__ import annotations

import argparse
import csv
import hashlib
import io
import urllib.request
from collections import defaultdict
from pathlib import Path

import openpyxl

SOURCE_STUDY_ID = "cyrille_etal_2025_martinique_gardens"
ARCHIPELAGO_ID = "lesser_antilles"
SOURCE_URL = "https://search-data.ubfc.fr/dl_data.php?file=601"
SOURCE_SHA256 = "e3f82dc81749d7c759dbb62fc2e40ceeff9382758a3114c63c57553d15c2327d"
SHEET = "Insects_Plants"
SYSTEMS = ("C1", "C2", "C3", "C4", "C5", "PG1", "PG2", "PG3", "PG4", "PG5")
PERIODS = tuple(f"P{i}" for i in range(1, 13))
EFFORT_MINUTES = 60.0
MIN_POSITIVE_BINS = 6
MIN_NONCONSTANT_SERIES = 3
OUTPUT_FIELDS = [
    "source_study_id",
    "archipelago_id",
    "system_id",
    "time_bin",
    "partner_id",
    "value",
    "effort",
]


def clean(value: object) -> str:
    return " ".join(str(value or "").strip().split())


def load_bytes(path: Path | None) -> bytes:
    if path is None:
        req = urllib.request.Request(
            SOURCE_URL,
            headers={"User-Agent": "izu-core-source-audit/1.0", "Accept": "*/*"},
        )
        with urllib.request.urlopen(req, timeout=120) as response:
            payload = response.read()
    else:
        payload = path.read_bytes()
    observed = hashlib.sha256(payload).hexdigest()
    if observed != SOURCE_SHA256:
        raise RuntimeError(f"Martinique source checksum drift: {observed}")
    return payload


def source_rows(payload: bytes) -> list[dict[str, object]]:
    wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise RuntimeError(f"missing sheet {SHEET!r}")
    ws = wb[SHEET]
    iterator = ws.iter_rows(values_only=True)
    headers = [clean(value) for value in next(iterator)]
    required = {"Period", "Site", "Insect_Best_ID", "Plant_Best_ID"}
    if not required.issubset(set(headers)):
        raise RuntimeError(f"Martinique schema drift: missing {sorted(required - set(headers))}")
    rows = [dict(zip(headers, values)) for values in iterator]
    wb.close()
    return rows


def reconstruct(payload: bytes) -> tuple[list[dict[str, object]], dict[str, dict[str, object]]]:
    events: dict[tuple[str, str, str], int] = defaultdict(int)
    observed_contexts: set[tuple[str, str]] = set()
    for row in source_rows(payload):
        system = clean(row.get("Site"))
        period = clean(row.get("Period"))
        if system not in SYSTEMS or period not in PERIODS:
            continue
        observed_contexts.add((system, period))
        partner = clean(row.get("Insect_Best_ID"))
        plant = clean(row.get("Plant_Best_ID"))
        # Mirrors the deposited source R code table(Plant_Best_ID, Insect_Best_ID).
        if not partner or not plant or partner.lower() in {"nan", "na"} or plant.lower() in {"nan", "na"}:
            continue
        events[(system, period, partner)] += 1

    expected_contexts = {(system, period) for system in SYSTEMS for period in PERIODS}
    missing = sorted(expected_contexts - observed_contexts)
    if missing:
        raise RuntimeError(f"Martinique source-native Site x Period coverage incomplete: {missing}")

    structure: dict[str, dict[str, object]] = {}
    system_series: dict[str, tuple[list[str], dict[str, list[int]]]] = {}
    for system in SYSTEMS:
        partners = sorted({partner for (s, _, partner) in events if s == system})
        if not partners:
            structure[system] = {
                "source_native_time_bins": len(PERIODS),
                "positive_interaction_bins": 0,
                "partner_count": 0,
                "nonconstant_partner_series": 0,
                "admitted": False,
                "reason": "FAIL_NO_IDENTIFIED_PARTNERS",
            }
            continue
        series_by_partner = {
            partner: [events.get((system, period, partner), 0) for period in PERIODS]
            for partner in partners
        }
        positive_bins = sum(
            any(series_by_partner[partner][i] > 0 for partner in partners)
            for i in range(len(PERIODS))
        )
        nonconstant = sum(len(set(series)) > 1 for series in series_by_partner.values())
        admitted = positive_bins >= MIN_POSITIVE_BINS and nonconstant >= MIN_NONCONSTANT_SERIES
        if positive_bins < MIN_POSITIVE_BINS:
            reason = "FAIL_LT6_POSITIVE_INTERACTION_BINS"
        elif nonconstant < MIN_NONCONSTANT_SERIES:
            reason = "FAIL_LT3_NONCONSTANT_PARTNER_SERIES"
        else:
            reason = "ADMIT_BEFORE_COORDINATES"
        structure[system] = {
            "source_native_time_bins": len(PERIODS),
            "positive_interaction_bins": positive_bins,
            "partner_count": len(partners),
            "nonconstant_partner_series": nonconstant,
            "admitted": admitted,
            "reason": reason,
        }
        if admitted:
            system_series[system] = (partners, series_by_partner)

    output: list[dict[str, object]] = []
    for system in SYSTEMS:
        if system not in system_series:
            continue
        partners, series_by_partner = system_series[system]
        for partner in partners:
            for period, count in zip(PERIODS, series_by_partner[partner]):
                if count <= 0:
                    continue
                output.append({
                    "source_study_id": SOURCE_STUDY_ID,
                    "archipelago_id": ARCHIPELAGO_ID,
                    "system_id": system,
                    "time_bin": period,
                    "partner_id": partner,
                    "value": float(count),
                    "effort": EFFORT_MINUTES,
                })
    if not output:
        raise RuntimeError("Martinique frozen structural gate admitted no systems")
    return output, structure


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=None)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--structure-out", type=Path, default=None)
    args = parser.parse_args()
    rows, structure = reconstruct(load_bytes(args.source))
    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    if args.structure_out is not None:
        import json
        args.structure_out.parent.mkdir(parents=True, exist_ok=True)
        args.structure_out.write_text(json.dumps(structure, indent=2, sort_keys=True) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
