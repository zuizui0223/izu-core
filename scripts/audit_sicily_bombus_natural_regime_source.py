from __future__ import annotations

import hashlib
import io
import json
import re
import urllib.request
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/results/chapter2_sicily_bombus_natural_regime_source_audit_20260914.json"
ARTICLE_DOI = "10.1007/s13592-025-01153-4"
ARTICLE_URL = "https://link.springer.com/article/10.1007/s13592-025-01153-4"
FILENAME = "13592_2025_1153_MOESM1_ESM.xlsx"
SOURCE_URLS = [
    "https://media.springernature.com/original/springer-static/esm/art%3A10.1007%2Fs13592-025-01153-4/MediaObjects/" + FILENAME,
    "https://media.springernature.com/original/springer-static/esm/art:10.1007/s13592-025-01153-4/MediaObjects/" + FILENAME,
    "https://static-content.springer.com/esm/art%3A10.1007%2Fs13592-025-01153-4/MediaObjects/" + FILENAME,
]
RAW_SITE_SHEETS = {
    "FIUMEDINISI transects": "Fiumedinisi-NAT",
    "BUTICARI transects": "Buticari-URB",
}


def fetch(url: str, timeout: int = 120) -> tuple[bytes, dict[str, str], str]:
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/140 Safari/537.36",
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*;q=0.8",
            "Referer": ARTICLE_URL,
        },
    )
    with urllib.request.urlopen(req, timeout=timeout) as response:
        return response.read(), {k.lower(): v for k, v in response.headers.items()}, response.geturl()


def clean(value: object) -> str:
    return " ".join(str(value or "").strip().split())


def norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", clean(value).casefold()).strip("_")


def first_header(rows: list[tuple]) -> tuple[int | None, list[str], list[tuple]]:
    for i, row in enumerate(rows):
        vals = [clean(v) for v in row]
        if len([v for v in vals if v]) >= 3:
            return i + 1, vals, rows[i + 1 :]
    return None, [], []


def role_columns(headers: list[str]) -> dict[str, list[str]]:
    roles = {"site": [], "date": [], "plant": [], "bombus": [], "count": [], "effort": []}
    for raw in headers:
        n = norm(raw)
        if not n:
            continue
        if any(token in n for token in ("site", "locality", "location", "place")):
            roles["site"].append(raw)
        if any(token in n for token in ("date", "day", "sampling_round", "round")):
            roles["date"].append(raw)
        if "plant" in n or "flower" in n:
            roles["plant"].append(raw)
        if any(token in n for token in ("bombus", "bumble", "bee_species", "species_bee")):
            roles["bombus"].append(raw)
        if any(token in n for token in ("count", "number", "abundance", "frequency", "n_ind")):
            roles["count"].append(raw)
        if any(token in n for token in ("effort", "duration", "hour", "minute", "transect")):
            roles["effort"].append(raw)
    return roles


def date_key(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value)
    return text[:10] if text else ""


def site_structure(headers: list[str], body: list[tuple], system_id: str) -> dict[str, object]:
    required = {"Date", "ID", "Bombus species", "Plant species"}
    if not required.issubset(set(headers)):
        return {"system_id": system_id, "status": "schema_fail", "missing": sorted(required - set(headers))}
    idx = {name: headers.index(name) for name in required}
    events: Counter[tuple[str, str]] = Counter()
    dates: set[str] = set()
    ids: set[str] = set()
    plants: set[str] = set()
    partners: set[str] = set()
    for row in body:
        raw_date = row[idx["Date"]] if idx["Date"] < len(row) else None
        raw_id = row[idx["ID"]] if idx["ID"] < len(row) else None
        raw_bee = row[idx["Bombus species"]] if idx["Bombus species"] < len(row) else None
        raw_plant = row[idx["Plant species"]] if idx["Plant species"] < len(row) else None
        d = date_key(raw_date)
        bee = clean(raw_bee)
        plant = clean(raw_plant)
        specimen = clean(raw_id)
        if not d or not bee or not plant:
            continue
        dates.add(d)
        partners.add(bee)
        plants.add(plant)
        if specimen:
            ids.add(specimen)
        events[(d, bee)] += 1

    ordered_dates = sorted(dates)
    per_year: dict[str, int] = Counter(d[:4] for d in ordered_dates)
    nonconstant = 0
    for bee in sorted(partners):
        series = [events.get((d, bee), 0) for d in ordered_dates]
        if len(set(series)) > 1:
            nonconstant += 1
    duplicate_ids = 0
    id_counts = Counter()
    for row in body:
        if idx["ID"] < len(row):
            specimen = clean(row[idx["ID"]])
            if specimen:
                id_counts[specimen] += 1
    duplicate_ids = sum(count > 1 for count in id_counts.values())
    expected_rounds_match = bool(per_year) and set(per_year).issubset({"2018", "2019"}) and all(v in {13, 14} for v in per_year.values()) and set(per_year) == {"2018", "2019"}
    return {
        "system_id": system_id,
        "status": "audited",
        "source_native_dates": len(ordered_dates),
        "dates_by_year": dict(sorted(per_year.items())),
        "date_min": ordered_dates[0] if ordered_dates else None,
        "date_max": ordered_dates[-1] if ordered_dates else None,
        "partner_count": len(partners),
        "plant_count": len(plants),
        "specimen_id_count": len(ids),
        "duplicate_specimen_ids": duplicate_ids,
        "nonconstant_partner_series": nonconstant,
        "published_13_to_14_rounds_per_site_year_matched": expected_rounds_match,
        "minimum_time_bins_pass": len(ordered_dates) >= 6,
        "minimum_nonconstant_partner_series_pass": nonconstant >= 3,
    }


def inspect_workbook(payload: bytes) -> dict[str, object]:
    wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    sheets = []
    site_structures = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        header_row, headers, body = first_header(rows)
        roles = role_columns(headers)
        record: dict[str, object] = {
            "sheet": ws.title,
            "nrows": ws.max_row,
            "ncols": ws.max_column,
            "header_row_1_based": header_row,
            "headers": headers,
            "roles": roles,
            "preview": [[str(v) if v is not None else None for v in row[:12]] for row in rows[:6]],
        }
        if ws.title in RAW_SITE_SHEETS:
            structure = site_structure(headers, body, RAW_SITE_SHEETS[ws.title])
            record["source_structure"] = structure
            site_structures.append(structure)
        sheets.append(record)
    wb.close()
    all_pass = bool(site_structures) and len(site_structures) == 2 and all(
        row.get("status") == "audited"
        and row.get("published_13_to_14_rounds_per_site_year_matched")
        and row.get("minimum_time_bins_pass")
        and row.get("minimum_nonconstant_partner_series_pass")
        for row in site_structures
    )
    return {
        "sheet_count": len(sheets),
        "sheets": sheets,
        "raw_site_sheets": list(RAW_SITE_SHEETS),
        "site_structures": site_structures,
        "all_structural_admission_checks_pass": all_pass,
    }


def main() -> None:
    state: dict[str, object] = {
        "schema_version": "1.2",
        "analysis": "chapter2_sicily_bombus_natural_regime_source_audit",
        "status": "source_structure_only_coordinates_not_opened",
        "source": {
            "article_doi": ARTICLE_DOI,
            "article_year": 2025,
            "island": "Sicily",
            "archipelago_id": "sicily",
            "official_supplement_urls": SOURCE_URLS,
            "published_design": {
                "fixed_sites": 2,
                "years": [2018, 2019],
                "sampling_window": "April-July",
                "sampling_rounds_per_site_per_year": "13-14",
                "transect_length_km": 1.0,
                "transect_duration_hours": 2.0,
                "same_collector": True,
                "source_note": "Each captured Bombus tube was marked with location and collection date, with associated plant identity; raw transect data are reported in DATAset.xlsx."
            }
        },
        "coordinates_opened": False,
    }

    attempts: list[dict[str, object]] = []
    accepted: bytes | None = None
    for url in SOURCE_URLS:
        try:
            payload, headers, resolved = fetch(url)
            attempt: dict[str, object] = {
                "url": url,
                "resolved_url": resolved,
                "bytes": len(payload),
                "sha256": hashlib.sha256(payload).hexdigest(),
                "content_type": headers.get("content-type"),
                "content_length_header": headers.get("content-length"),
                "content_disposition": headers.get("content-disposition"),
                "prefix_hex": payload[:32].hex(),
                "xlsx_signature": payload.startswith(b"PK\x03\x04"),
            }
            if payload.startswith(b"PK\x03\x04"):
                attempt["status"] = "accepted_xlsx_payload"
                attempts.append(attempt)
                accepted = payload
                break
            attempt["status"] = "non_xlsx_payload"
            attempt["prefix_text"] = payload[:300].decode("utf-8", errors="replace")
            attempts.append(attempt)
        except Exception as exc:
            attempts.append({"url": url, "status": "request_failed", "error": repr(exc)})
    state["transport_attempts"] = attempts

    if accepted is None:
        state["decision"] = "TRANSPORT_BLOCKED_OFFICIAL_SPRINGER_PATHS_RETURN_NO_XLSX"
    else:
        state["download"] = {
            "status": "recovered_valid_xlsx",
            "bytes": len(accepted),
            "sha256": hashlib.sha256(accepted).hexdigest(),
        }
        workbook = inspect_workbook(accepted)
        state["workbook"] = workbook
        state["decision"] = (
            "SOURCE_STRUCTURE_PASS_READY_TO_FREEZE_ADAPTER"
            if workbook["all_structural_admission_checks_pass"]
            else "SOURCE_RECOVERED_BUT_STRUCTURAL_ADMISSION_CHECK_FAILED"
        )

    state["biological_negative"] = False
    state["claim_boundary"] = (
        "This audit uses only lawful source-equivalent Springer supplement transport, source-native site sheets, dates, specimen IDs, Bombus labels, plant labels, and the published equal two-hour transect design. "
        "The published 13-14 rounds per site-year are checked before any natural-regime coordinate is opened. No D1, phi, response, determinant, or route value is calculated."
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(state, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({
        "decision": state["decision"],
        "download": state.get("download"),
        "site_structures": (state.get("workbook") or {}).get("site_structures", []),
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
