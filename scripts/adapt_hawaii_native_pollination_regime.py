from __future__ import annotations

import argparse
import csv
import hashlib
import math
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

SOURCE_ID = "aslan_etal_2019_hawaii_native_pollination"
ARCHIPELAGO_ID = "hawaiian_islands"
SYSTEM_ID = "pohakuloa_high_elevation_dryland_pollination_community"
OBSERVATION_SHEETS = (
    "ARGGLA",
    "BIDMEN",
    "DUBLIN",
    "HAPHAP",
    "SIDFAL",
    "SILLAN",
    "STEANG",
    "TETARE",
)
WORKBOOK_SHA256 = "2b0ff40226b2a6d511a111ead8a00660532de3d799aed217e4dc30f00c2b3c27"
MISSING_SCAN_LABELS = {"", "N", "NONE", "UNKNOWN", "?", "N/A", "NA"}
OUTPUT_COLUMNS = (
    "source_study_id",
    "archipelago_id",
    "system_id",
    "time_bin",
    "partner_id",
    "value",
    "effort",
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _partner_id(value: object) -> str:
    return _clean(value).upper()


def _date_id(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = _clean(value)
    if not text:
        return ""
    try:
        return datetime.fromisoformat(text).date().isoformat()
    except ValueError:
        return text


def _event_key(sheet: str, row: dict[str, object]) -> tuple[str, ...] | None:
    date = _date_id(row.get("Date"))
    start = _clean(row.get("Start Time"))
    observer = _clean(row.get("Observer")).upper()
    block = _clean(row.get("Scan block time start"))
    if not all((date, start, observer, block)):
        return None
    return (
        sheet,
        _clean(row.get("Site")).upper(),
        date,
        start,
        observer,
        block,
    )


def _rows(worksheet) -> list[dict[str, object]]:
    values = list(worksheet.iter_rows(values_only=True))
    headers = [_clean(value) for value in values[0]]
    return [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]


def adapt_workbook(
    workbook_path: Path,
    *,
    expected_sha256: str = WORKBOOK_SHA256,
) -> tuple[list[dict[str, object]], dict[str, object]]:
    observed_sha = _sha256(workbook_path)
    if observed_sha != expected_sha256:
        raise ValueError(f"unexpected Hawaii workbook SHA256: {observed_sha}")

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    missing_sheets = sorted(set(OBSERVATION_SHEETS) - set(workbook.sheetnames))
    if missing_sheets:
        workbook.close()
        raise ValueError(f"missing observation sheets: {missing_sheets}")

    events_by_date: dict[str, set[tuple[str, ...]]] = defaultdict(set)
    counts: dict[tuple[str, str], float] = defaultdict(float)
    event_partner_values: dict[tuple[tuple[str, ...], str], float] = {}
    raw_partner_labels: set[str] = set()

    for sheet in OBSERVATION_SHEETS:
        for row in _rows(workbook[sheet]):
            event = _event_key(sheet, row)
            if event is None:
                continue
            date = event[2]
            events_by_date[date].add(event)
            partner = _partner_id(row.get("Scan visitor spp"))
            if partner in MISSING_SCAN_LABELS:
                continue
            value = row.get("Scan # Inds")
            if value is None or _clean(value).upper() in MISSING_SCAN_LABELS:
                # Retain the scan event in effort but do not impute an unquantified numerator.
                continue
            if not isinstance(value, (int, float)):
                raise ValueError(f"nonnumeric scan count for {sheet} {date} {partner}: {value!r}")
            value = float(value)
            if value < 0 or not math.isfinite(value):
                raise ValueError(f"invalid scan count for {sheet} {date} {partner}: {value!r}")
            key = (event, partner)
            if key in event_partner_values and event_partner_values[key] != value:
                raise ValueError(
                    f"conflicting duplicate scan record for {key}: "
                    f"{event_partner_values[key]} vs {value}"
                )
            event_partner_values[key] = value
            raw_partner_labels.add(partner)

    workbook.close()

    for (event, partner), value in event_partner_values.items():
        counts[(event[2], partner)] += value

    dates = sorted(events_by_date)
    partners = sorted(raw_partner_labels)
    if not dates or not partners:
        raise ValueError("no usable Hawaii scan data")

    output: list[dict[str, object]] = []
    for date in dates:
        effort = len(events_by_date[date])
        if effort <= 0:
            raise ValueError(f"nonpositive scan effort for {date}")
        for partner in partners:
            output.append(
                {
                    "source_study_id": SOURCE_ID,
                    "archipelago_id": ARCHIPELAGO_ID,
                    "system_id": SYSTEM_ID,
                    "time_bin": date,
                    "partner_id": partner,
                    "value": counts.get((date, partner), 0.0),
                    "effort": effort,
                }
            )

    standardized: dict[str, list[float]] = {partner: [] for partner in partners}
    nonempty_dates = 0
    for date in dates:
        effort = len(events_by_date[date])
        date_total = 0.0
        for partner in partners:
            value = counts.get((date, partner), 0.0) / effort
            standardized[partner].append(value)
            date_total += value
        nonempty_dates += int(date_total > 0)
    nonconstant = sum(1 for values in standardized.values() if max(values) > min(values))

    diagnostics = {
        "source_study_id": SOURCE_ID,
        "archipelago_id": ARCHIPELAGO_ID,
        "system_id": SYSTEM_ID,
        "workbook_sha256": observed_sha,
        "observation_sheets": list(OBSERVATION_SHEETS),
        "time_bins": len(dates),
        "first_time_bin": dates[0],
        "last_time_bin": dates[-1],
        "unique_scan_events": sum(len(value) for value in events_by_date.values()),
        "scan_events_per_time_bin_min": min(len(value) for value in events_by_date.values()),
        "scan_events_per_time_bin_max": max(len(value) for value in events_by_date.values()),
        "raw_partner_labels": len(partners),
        "nonconstant_partner_series": nonconstant,
        "nonempty_time_bins": nonempty_dates,
        "system_count": 1,
        "coordinate_values_opened": False,
    }
    return output, diagnostics


def write_canonical_csv(rows: list[dict[str, object]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    rows, diagnostics = adapt_workbook(args.workbook)
    write_canonical_csv(rows, args.out)
    print(diagnostics)


if __name__ == "__main__":
    main()
