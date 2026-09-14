from __future__ import annotations

import argparse
import csv
import hashlib
import math
from collections import defaultdict
from pathlib import Path

import openpyxl

SOURCE_ID = "lazaro_etal_2022_mallorca_stability"
ARCHIPELAGO_ID = "balearic_islands"
DATA_SHA256 = "86ea6ae76539cd525509096e49a00a15f9cd0b3faeb10019cb92109b7ae1b1f3"
MINIMUM_TIME_BINS = 6
OUTPUT_COLUMNS = (
    "source_study_id",
    "archipelago_id",
    "system_id",
    "time_bin",
    "partner_id",
    "value",
    "effort",
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _clean(value: object) -> str:
    return str(value or "").strip()


def _sheet_rows(workbook, sheet: str) -> list[dict[str, object]]:
    values = list(workbook[sheet].iter_rows(values_only=True))
    headers = [_clean(value) for value in values[0]]
    return [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]


def split_pair_identities(
    pairs: set[str], insect_vocabulary: set[str], plant_vocabulary: set[str]
) -> dict[str, tuple[str, str]]:
    insects = set(insect_vocabulary)
    plants = set(plant_vocabulary)
    remaining = set(pairs)
    resolved: dict[str, tuple[str, str]] = {}

    while remaining:
        changed = False
        for pair in sorted(remaining):
            plant_matches = [name for name in plants if pair.endswith(name) and len(pair) > len(name)]
            if plant_matches:
                longest = max(map(len, plant_matches))
                plant_matches = [name for name in plant_matches if len(name) == longest]
            if len(plant_matches) == 1:
                plant = plant_matches[0]
                insect = pair[: -len(plant)].strip()
                if not insect:
                    raise ValueError(f"empty insect identity after plant-suffix split: {pair!r}")
                resolved[pair] = (insect, plant)
                insects.add(insect)
                remaining.remove(pair)
                changed = True
                continue

            insect_matches = [name for name in insects if pair.startswith(name) and len(pair) > len(name)]
            if insect_matches:
                longest = max(map(len, insect_matches))
                insect_matches = [name for name in insect_matches if len(name) == longest]
            if len(insect_matches) == 1:
                insect = insect_matches[0]
                plant = pair[len(insect) :].strip()
                if not plant:
                    raise ValueError(f"empty plant identity after insect-prefix split: {pair!r}")
                resolved[pair] = (insect, plant)
                plants.add(plant)
                remaining.remove(pair)
                changed = True
                continue

            if len(plant_matches) > 1 or len(insect_matches) > 1:
                raise ValueError(f"ambiguous source pair identity: {pair!r}")

        if not changed:
            raise ValueError(f"unresolved source pair identities: {sorted(remaining)[:20]}")

    return resolved


def adapt_workbook(
    workbook_path: Path,
    *,
    expected_sha256: str = DATA_SHA256,
) -> tuple[list[dict[str, object]], dict[str, object]]:
    observed_sha = _sha256(workbook_path)
    if observed_sha != expected_sha256:
        raise ValueError(f"unexpected Mallorca workbook SHA256: {observed_sha}")

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    required_sheets = {"IndividualPlants_Abundance", "Insect_Abundance", "Interactions_Abundance"}
    missing = sorted(required_sheets - set(workbook.sheetnames))
    if missing:
        workbook.close()
        raise ValueError(f"missing source sheets: {missing}")

    plant_rows = _sheet_rows(workbook, "IndividualPlants_Abundance")
    insect_rows = _sheet_rows(workbook, "Insect_Abundance")
    interaction_rows = _sheet_rows(workbook, "Interactions_Abundance")
    workbook.close()

    plants = {_clean(row["PlSpecies"]) for row in plant_rows if _clean(row.get("PlSpecies"))}
    insects = {_clean(row["InSpecies"]) for row in insect_rows if _clean(row.get("InSpecies"))}
    pairs = {_clean(row["InsectPlantPair"]) for row in interaction_rows if _clean(row.get("InsectPlantPair"))}
    pair_split = split_pair_identities(pairs, insects, plants)

    rounds_by_locality: dict[str, set[int]] = defaultdict(set)
    partners_by_locality: dict[str, set[str]] = defaultdict(set)
    counts: dict[tuple[str, int, str], float] = defaultdict(float)
    for row in interaction_rows:
        locality = _clean(row.get("Locality"))
        pair = _clean(row.get("InsectPlantPair"))
        if not locality or not pair:
            continue
        round_value = row.get("Round")
        abundance = row.get("Abundance")
        if not isinstance(round_value, (int, float)) or int(round_value) != float(round_value):
            raise ValueError(f"invalid Round for {locality}: {round_value!r}")
        if not isinstance(abundance, (int, float)):
            raise ValueError(f"nonnumeric Abundance for {locality}: {abundance!r}")
        abundance = float(abundance)
        if abundance < 0 or not math.isfinite(abundance):
            raise ValueError(f"invalid Abundance for {locality}: {abundance!r}")
        round_id = int(round_value)
        insect, _plant = pair_split[pair]
        rounds_by_locality[locality].add(round_id)
        partners_by_locality[locality].add(insect)
        counts[(locality, round_id, insect)] += abundance

    admitted = sorted(
        locality for locality, rounds in rounds_by_locality.items() if len(rounds) >= MINIMUM_TIME_BINS
    )
    excluded = sorted(set(rounds_by_locality) - set(admitted))
    output: list[dict[str, object]] = []
    structural = []
    for locality in admitted:
        rounds = sorted(rounds_by_locality[locality])
        partners = sorted(partners_by_locality[locality])
        series = {partner: [] for partner in partners}
        for round_id in rounds:
            for partner in partners:
                value = counts.get((locality, round_id, partner), 0.0)
                series[partner].append(value)
                output.append(
                    {
                        "source_study_id": SOURCE_ID,
                        "archipelago_id": ARCHIPELAGO_ID,
                        "system_id": locality,
                        "time_bin": f"round_{round_id}",
                        "partner_id": partner,
                        "value": value,
                        "effort": 1.0,
                    }
                )
        nonconstant = sum(1 for values in series.values() if max(values) > min(values))
        structural.append(
            {
                "system_id": locality,
                "time_bins": len(rounds),
                "rounds": rounds,
                "partner_count": len(partners),
                "nonconstant_partner_series": nonconstant,
            }
        )

    diagnostics = {
        "source_study_id": SOURCE_ID,
        "archipelago_id": ARCHIPELAGO_ID,
        "workbook_sha256": observed_sha,
        "source_pair_identities": len(pairs),
        "resolved_pair_identities": len(pair_split),
        "source_localities": len(rounds_by_locality),
        "admitted_systems": len(admitted),
        "excluded_systems": [
            {"system_id": locality, "time_bins": len(rounds_by_locality[locality])}
            for locality in excluded
        ],
        "system_structure": structural,
        "coordinate_values_opened": False,
    }
    return output, diagnostics


def write_canonical_csv(rows: list[dict[str, object]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    rows, diagnostics = adapt_workbook(args.workbook)
    write_canonical_csv(rows, args.out)
    print(diagnostics)


if __name__ == "__main__":
    main()
