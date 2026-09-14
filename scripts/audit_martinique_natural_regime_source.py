from __future__ import annotations

import hashlib
import json
import math
import re
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/results/chapter2_martinique_natural_regime_source_audit_20260914.json"
RAW = ROOT / "data/external/martinique_natural_regime"
USER_AGENT = "izu-core-source-audit/1.0"

FILES = {
    "Sampling_data.xlsx": "https://search-data.ubfc.fr/dl_data.php?file=601",
    "Plant_insect_interactions_former_names.xlsx": "https://search-data.ubfc.fr/dl_data.php?file=597",
    "Interaction_turnover.R": "https://search-data.ubfc.fr/dl_data.php?file=589",
    "Sampling_completeness.R": "https://search-data.ubfc.fr/dl_data.php?file=600",
    "README.docx": "https://search-data.ubfc.fr/dl_data.php?file=599",
}


def download(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT, "Accept": "*/*"})
    with urllib.request.urlopen(req, timeout=120) as response:
        return response.read()


def norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def first_header_and_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    for idx, row in enumerate(rows):
        nonempty = [v for v in row if v not in (None, "")]
        if len(nonempty) >= 3:
            headers = [str(v or "").strip() for v in row]
            return idx + 1, headers, rows[idx + 1 :]
    return None, [], []


def role_columns(headers: list[str]) -> dict[str, list[str]]:
    out = {"site": [], "time": [], "plant": [], "partner": [], "interaction": [], "effort": []}
    for raw in headers:
        n = norm(raw)
        if not n:
            continue
        if any(k in n for k in ("site", "garden", "jardin")):
            out["site"].append(raw)
        if any(k in n for k in ("date", "month", "mois", "period", "periode", "sampling")):
            out["time"].append(raw)
        if "plant" in n or "plante" in n:
            out["plant"].append(raw)
        if any(k in n for k in ("insect", "pollinator", "visitor", "visiteur")):
            out["partner"].append(raw)
        if any(k in n for k in ("interaction", "occurrence", "visit", "visite", "count", "nombre", "abundance")):
            out["interaction"].append(raw)
        if any(k in n for k in ("duration", "minute", "effort", "transect", "length", "distance")):
            out["effort"].append(raw)
    return out


def value_set(headers: list[str], rows: list[tuple], column: str) -> set[str]:
    try:
        i = headers.index(column)
    except ValueError:
        return set()
    return {str(row[i]).strip() for row in rows if i < len(row) and row[i] not in (None, "")}


def numeric_presence(headers: list[str], rows: list[tuple], column: str) -> dict:
    try:
        i = headers.index(column)
    except ValueError:
        return {"numeric": 0, "nonmissing": 0}
    numeric = 0
    nonmissing = 0
    for row in rows:
        if i >= len(row) or row[i] in (None, ""):
            continue
        nonmissing += 1
        try:
            x = float(row[i])
            if math.isfinite(x):
                numeric += 1
        except (TypeError, ValueError):
            pass
    return {"numeric": numeric, "nonmissing": nonmissing}


def inventory_workbook(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        header_row, headers, rows = first_header_and_rows(ws)
        roles = role_columns(headers)
        sheet = {
            "sheet": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "header_row_1_based": header_row,
            "headers": headers,
            "roles": roles,
        }
        for role in ("site", "time", "plant", "partner"):
            sheet[f"{role}_cardinalities"] = {
                col: len(value_set(headers, rows, col)) for col in roles[role]
            }
        sheet["interaction_numeric_presence"] = {
            col: numeric_presence(headers, rows, col) for col in roles["interaction"]
        }
        sheets.append(sheet)
    wb.close()
    return sheets


def inventory_r_code(payload: bytes) -> dict:
    text = payload.decode("utf-8-sig", errors="replace")
    indicators = []
    patterns = (
        "sampling_data", "insects_plants", "insect_best_id", "plant_best_id",
        "num_sp", "group_by", "summarise", "summarize", "count(", "n()",
        "pivot_wider", "network", "interaction", "period", "site",
    )
    for line_no, line in enumerate(text.splitlines(), start=1):
        lower = line.lower()
        if any(pattern in lower for pattern in patterns):
            indicators.append({"line": line_no, "text": line[:500]})
    return {
        "line_count": len(text.splitlines()),
        "structural_indicator_lines": indicators[:300],
        "coordinates_opened": False,
    }


def main() -> None:
    RAW.mkdir(parents=True, exist_ok=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    files = {}
    for name, url in FILES.items():
        try:
            payload = download(url)
            path = RAW / name
            path.write_bytes(payload)
            record = {
                "status": "recovered",
                "bytes": len(payload),
                "sha256": hashlib.sha256(payload).hexdigest(),
                "url": url,
            }
            if name.endswith(".xlsx"):
                record["workbook"] = inventory_workbook(path)
            elif name.endswith(".R"):
                record["r_code"] = inventory_r_code(payload)
            files[name] = record
        except Exception as exc:
            files[name] = {"status": "failed", "url": url, "error": repr(exc)}

    workbook_records = [r for n, r in files.items() if n.endswith(".xlsx") and r.get("status") == "recovered"]
    candidate_sheets = []
    for record in workbook_records:
        for sheet in record.get("workbook", []):
            roles = sheet["roles"]
            if roles["site"] and roles["time"] and roles["partner"]:
                candidate_sheets.append({
                    "sheet": sheet["sheet"],
                    "headers": sheet["headers"],
                    "roles": roles,
                    "site_cardinalities": sheet["site_cardinalities"],
                    "time_cardinalities": sheet["time_cardinalities"],
                    "partner_cardinalities": sheet["partner_cardinalities"],
                    "interaction_numeric_presence": sheet["interaction_numeric_presence"],
                })

    r_code_recovered = [
        name for name, record in files.items()
        if name.endswith(".R") and record.get("status") == "recovered"
    ]
    result = {
        "schema_version": "1.1",
        "analysis": "chapter2_martinique_natural_regime_source_audit",
        "status": "source_structure_only_coordinates_not_opened",
        "source": {
            "dataset_doi": "10.25666/DATAUBFC-2025-03-28",
            "island": "Martinique",
            "archipelago": "Lesser Antilles",
            "public_record_created": "2025-03-28",
            "sampling_period": "2022-10-01 to 2023-09-12",
            "design": "10 garden sites sampled monthly for 12 months",
            "insect_effort": "60 minutes along 150 m per site per month by the same observer",
        },
        "files": files,
        "candidate_interaction_sheets": candidate_sheets,
        "recovered_source_r_scripts": r_code_recovered,
        "gate_design_checks_from_public_metadata": {
            "public_before_frozen_cutoff": True,
            "stable_site_identity_expected": True,
            "source_native_months_expected": 12,
            "minimum_time_bins_design_pass": True,
            "effort_equal_by_design_for_insect_interactions": True,
            "coordinates_opened": False,
        },
        "decision": (
            "RAW_SCHEMA_AND_SOURCE_CODE_RECOVERED_NEXT_FREEZE_ADAPTER"
            if candidate_sheets and r_code_recovered
            else "RAW_SCHEMA_RECOVERED_SOURCE_CODE_INCOMPLETE"
            if candidate_sheets
            else "BLOCKED_OR_SCHEMA_NOT_YET_IDENTIFIED"
        ),
        "claim_boundary": (
            "This audit inspects only public file recovery, workbook headers, identifier cardinalities and source-code structural operations. "
            "It does not aggregate interaction values or compute D1, phi, route thresholds, or outcome variables. "
            "Bird interactions are not promoted to the monthly weighted primary matrix because the source metadata says they were reduced to binary site-by-bi-period pairs."
        ),
    }
    OUT.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
